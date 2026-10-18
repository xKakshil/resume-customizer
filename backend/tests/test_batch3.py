"""Layout ATS check, grounding drift, one-page fit, export read-back, interview prep."""

import io

import pytest

from app.models.schemas import TailoredResume, TailorRequest
from app.services.export import pdf_page_count, resume_pdf
from app.services.layout_check import docx_layout_checks, pdf_layout_checks
from app.services.tailor import apply_guardrails, grounding
from tests.fixtures import SAMPLE_JD, make_docx, make_pdf, resume_text
from tests.test_tailor import make


def _status(checks, cid):
    return next(c["status"] for c in checks if c["id"] == cid)


def _two_column_pdf(with_image=False, header_contact=False) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    if header_contact:
        c.drawString(40, 830, "jane@example.com | +91 98765 43210")
    y = 780
    for i in range(30):
        c.drawString(40, y, f"Skill line {i} with Python and SQL")
        c.drawString(320, y, f"Built API number {i} for the payments team")
        y -= 20
    if with_image:
        from reportlab.lib.utils import ImageReader
        from PIL import Image
        img = io.BytesIO()
        Image.new("RGB", (20, 20), "red").save(img, "PNG")
        img.seek(0)
        c.drawImage(ImageReader(img), 500, 60, 20, 20)
    c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_layout_single_column_passes():
    checks = pdf_layout_checks(make_pdf())
    assert all(c["status"] == "pass" for c in checks), checks


def test_pdf_layout_detects_columns_images_and_header_contact():
    checks = pdf_layout_checks(_two_column_pdf(with_image=True, header_contact=True))
    assert _status(checks, "columns") == "fail"
    assert _status(checks, "images") == "warn"
    assert _status(checks, "margins") == "warn"
    assert all(c["tip"] for c in checks if c["status"] != "pass")


def test_our_exported_pdf_passes_layout_checks():
    r = TailoredResume.model_validate({"candidate": {"name": "A", "contact": {"email": "a@b.co"}},
                                       "experience": [{"title": "Intern", "company": "X", "dates": "Jan 2025 – Jun 2025",
                                                       "bullets": [{"original": "Built APIs in Python serving 10k users daily.", "tailored": ""}] * 6}]})
    assert all(c["status"] == "pass" for c in pdf_layout_checks(resume_pdf(r)))


def test_docx_layout_checks():
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "jane@example.com"
    d.add_paragraph("Experience")
    buf = io.BytesIO()
    d.save(buf)
    checks = docx_layout_checks(buf.getvalue())
    assert _status(checks, "margins") == "warn" and _status(checks, "tables") == "pass"
    assert _status(docx_layout_checks(make_docx()), "tables") == "warn"


def test_parse_includes_layout_and_warns_on_columns():
    from app.services.parsing import parse_resume
    lines = ["Jane Doe", "jane@example.com", "EXPERIENCE", "EDUCATION"]
    p = parse_resume(_two_column_pdf(), "two.pdf")
    assert any(c["id"] == "columns" for c in p.layout)
    assert any("columns" in w for w in p.warnings)
    assert lines  # noqa


# --------------------------------------------------------------------------- #
def test_grounding_score():
    o = "Worked on the payments dashboard using React and TypeScript used by 12,000 merchants."
    assert grounding(o, "Built the payments dashboard in React and TypeScript for 12,000 merchants") > 0.6
    assert grounding(o, "Led a company-wide platform migration to microservices") < 0.3


def test_low_grounding_bullets_are_flagged():
    r = make()
    r.experience[0].bullets[0].tailored = "Architected an event-driven microservices platform handling enterprise workloads"
    apply_guardrails(r, TailorRequest(job_description=SAMPLE_JD, resume_text=resume_text()), resume_text(), [])
    assert "low_grounding" in r.experience[0].bullets[0].flags


# --------------------------------------------------------------------------- #
def _long_resume(n_bullets: int) -> TailoredResume:
    bullet = {"original": "x", "tailored": "Built and maintained REST APIs in Python and PostgreSQL, cutting response times by 35% for 12,000 merchants."}
    return TailoredResume.model_validate({
        "candidate": {"name": "Aarav Sharma", "contact": {"email": "a@example.com"}},
        "summary": {"tailored": "Backend engineer " * 12},
        "skills": {"groups": [{"category": "Languages", "items": ["Python", "TypeScript", "SQL"]}]},
        "experience": [{"title": f"Engineer {i}", "company": "Acme", "dates": "2024", "bullets": [bullet] * 4} for i in range(n_bullets // 4)],
        "education": [{"degree": "B.Tech", "institution": "IIT (BHU)"}],
    })


def test_one_page_fit_compacts_slight_overflow():
    from app.services.export import _resume_pdf
    r = None
    for n in range(16, 60, 4):  # find a resume that just overflows one page at normal density
        cand = _long_resume(n)
        _, normal = _resume_pdf(cand, compact=False)
        _, tight = _resume_pdf(cand, compact=True)
        if normal == 2 and tight == 1:
            r = cand
            break
    assert r is not None, "no slight-overflow case found"
    assert pdf_page_count(resume_pdf(r)) == 1


def test_export_check_round_trip(client):
    from tests.test_api import payload
    resume = client.post("/api/tailor", json=payload()).json()["resume"]
    res = client.post("/api/export/check", json={"format": "pdf", "resume": resume}).json()
    assert res["bullets_total"] > 0 and res["bullets_found"] == res["bullets_total"]
    assert set(res["sections_found"]) == set(res["sections_expected"]) and res["contact_found"] and res["pages"] >= 1


# --------------------------------------------------------------------------- #
def test_interview_prep(client):
    from tests.test_api import payload
    r = client.post("/api/interview-prep", json={**payload(), "count": 6})
    assert r.status_code == 200, r.text
    body = r.json()
    assert 0 < len(body["questions"]) <= 6 and body["role_fit"] and body["questions_to_ask"]
    assert {q["category"] for q in body["questions"]} <= {"behavioral", "technical", "project", "role"}
    assert all(q["answer_points"] for q in body["questions"])


def test_interview_prep_bad_output(client, monkeypatch):
    from app.api import interview
    from app.services.llm.factory import ProviderChain
    from app.services.llm.mock import MockProvider

    class Bad(MockProvider):
        async def complete(self, *a, **k):
            return "not json"
    monkeypatch.setattr(interview, "get_chain", lambda: ProviderChain([Bad()]))
    from tests.test_api import payload
    r = client.post("/api/interview-prep", json=payload())
    assert r.status_code == 502 and r.json()["error"]["code"] == "llm_bad_output"


def test_history_attach_interview_prep(client):
    from tests.test_api import payload
    tok = client.post("/api/auth/register", json={"email": "prep@example.com", "password": "supersecret1"}).json()["token"]
    h = {"Authorization": f"Bearer {tok}"}
    rid = client.post("/api/tailor", json=payload(), headers=h).json()["id"]
    assert client.put(f"/api/history/{rid}/interview-prep", json={"prep": {"questions": []}}, headers=h).status_code == 200
    assert client.get(f"/api/history/{rid}", headers=h).json()["result"]["interview_prep"] == {"questions": []}
