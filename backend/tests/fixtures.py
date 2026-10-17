"""Builders for realistic test resumes (PDF/DOCX) and a sample JD."""

from __future__ import annotations

import io

RESUME_LINES = [
    "Aarav Sharma",
    "aarav.sharma@example.com | +91 98765 43210 | Bengaluru, India | github.com/aarav-dev | linkedin.com/in/aarav-sharma",
    "",
    "SUMMARY",
    "Software engineer with experience building web applications and APIs. Enjoys solving problems and learning new technologies.",
    "",
    "EXPERIENCE",
    "Software Engineer Intern — Finlytics, Bengaluru | Jan 2025 – Jun 2025",
    "• Worked on the payments dashboard using React and TypeScript used by 12,000 merchants.",
    "• Responsible for building REST APIs in Node.js and Express backed by Postgres, cutting p95 latency by 35%.",
    "• Wrote unit tests with Jest and set up GitHub Actions CI for the frontend repo.",
    "• Helped migrate cron jobs to a Redis-backed queue, reducing failed jobs by 60%.",
    "",
    "Frontend Developer Intern — Opined | May 2024 – Jul 2024",
    "• Built reusable React components and improved Lighthouse performance score from 62 to 91.",
    "• Collaborated with designers in Figma to ship 8 responsive pages.",
    "",
    "PROJECTS",
    "QuickLink — URL shortener | TypeScript, Express, MongoDB, Redis",
    "• Built a URL shortener with click analytics handling 1,000 requests per second in load tests.",
    "• Deployed with Docker on Render with Redis caching for hot links.",
    "",
    "EDUCATION",
    "B.Tech, Electronics Engineering — IIT (BHU) Varanasi | 2022 – 2026",
    "CGPA: 8.1/10",
    "",
    "SKILLS",
    "Languages: Python, TypeScript, JavaScript, C++, SQL",
    "Frameworks: React, Next.js, Node.js, Express, FastAPI",
    "Tools: Git, Docker, PostgreSQL, MongoDB, Redis, Jest, Figma",
    "",
    "ACHIEVEMENTS",
    "• Codeforces Expert (max rating 1650).",
]

SAMPLE_JD = """Backend Engineer (SDE-1) — Acme Payments

About the role
Acme Payments is hiring a Backend Engineer to build and scale the APIs that power payouts for 50,000 businesses.

Responsibilities
- Design, build and maintain RESTful APIs and microservices in Python (FastAPI) or Node.js
- Own services end-to-end: design, testing, deployment and monitoring
- Improve performance and reliability of PostgreSQL-backed systems; use Redis for caching
- Write unit and integration tests; participate in code reviews

Requirements (must have)
- Strong knowledge of Python or TypeScript, data structures and algorithms
- Experience with PostgreSQL, Redis and REST API design
- Familiarity with Docker and CI/CD (GitHub Actions)

Nice to have
- Kubernetes, AWS, Kafka
- Experience with payments or fintech
- SOC2 awareness
"""


def resume_text() -> str:
    return "\n".join(RESUME_LINES)


def make_pdf(lines: list[str] | None = None, *, encrypt: str | None = None, pages: int = 1, blank: bool = False) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, encrypt=encrypt)
    lines = lines or RESUME_LINES
    for _ in range(pages):
        if blank:
            c.rect(50, 50, 200, 200, fill=1)  # an "image" with no text, like a scan
        else:
            y = 800
            for line in lines:
                c.setFont("Helvetica", 10)
                c.drawString(40, y, line.replace("—", "-").replace("–", "-"))
                y -= 14
            c.linkURL("https://github.com/aarav-dev", (40, 770, 200, 790))
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(lines: list[str] | None = None, *, with_table: bool = True) -> bytes:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    d = docx.Document()
    for line in lines or RESUME_LINES:
        if line.startswith("• "):
            d.add_paragraph(line[2:], style="List Bullet")
        else:
            d.add_paragraph(line)
    if with_table:
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Certifications"
        t.cell(0, 1).text = "AWS Cloud Practitioner"
    d.part.relate_to("https://portfolio.aarav.dev", RT.HYPERLINK, is_external=True)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
