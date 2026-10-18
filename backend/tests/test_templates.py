"""LaTeX input, template matching (original DOCX/TeX edited in place), page setup,
inline previews and supporting documents."""

import io
from pathlib import Path
from urllib.parse import unquote

import pytest

from app.models.schemas import LayoutSettings, TailoredResume
from app.services.latex_source import apply_to_tex, bullet_spans, compile_tex, latex_engine, tex_to_text
from tests.fixtures import SAMPLE_JD, make_docx, make_pdf

JAKE = (Path(__file__).parent / "data" / "jake.tex").read_text()


def _payload(text):
    return {"job_description": SAMPLE_JD, "resume_text": text}


# --------------------------------------------------------------------------- #
def test_tex_to_text_reads_jake_template():
    text, links = tex_to_text(JAKE)
    assert text.startswith("Aarav Sharma")
    assert "• Worked on the payments dashboard using React and TypeScript used by 12,000 merchants." in text
    assert "cutting p95 latency by 35%." in text
    assert "Languages: Python, TypeScript" in text
    assert "\\" not in text and "{" not in text
    assert "https://github.com/aarav-dev" in links
    assert len(bullet_spans(JAKE)) == 9  # 8 bullets + the skills \item


def test_apply_to_tex_swaps_reorders_and_preserves_the_rest():
    B = lambda o, t: {"original": o, "tailored": t}
    r = TailoredResume.model_validate({"summary": {"tailored": "Builds APIs & tests 100%."}, "experience": [{"bullets": [
        B("Helped migrate cron jobs to a Redis-backed queue, reducing failed jobs by 60%.", "Migrated cron jobs to a Redis-backed queue, cutting failed jobs by 60%."),
        B("Worked on the payments dashboard using React and TypeScript used by 12,000 merchants.", "Worked on the payments dashboard using React and TypeScript used by 12,000 merchants."),
    ]}]})
    out, rep = apply_to_tex(JAKE, r, show_summary=False)
    assert rep["summary"] == "removed" and "\\section{Summary}" not in out
    assert "\\resumeItem{Migrated cron jobs to a Redis-backed queue, cutting failed jobs by 60\\%.}" in out
    # The unchanged bullet moved to the second slot and kept its inline \textbf formatting.
    assert out.index("cutting failed jobs") < out.index("payments dashboard using \\textbf{React}")
    assert out.split("\\begin{document}")[0] == JAKE.split("\\begin{document}")[0]  # preamble untouched
    out2, rep2 = apply_to_tex(JAKE, r, show_summary=True)
    assert rep2["summary"] == "replaced" and "Builds APIs \\& tests 100\\%." in out2


@pytest.mark.skipif(latex_engine() is None, reason="no LaTeX engine installed")
def test_compile_is_sandboxed():
    assert compile_tex(JAKE) is not None
    for evil in ("\\input{/etc/passwd}", "\\openin5=/etc/passwd", "^^5cinput{/etc/passwd}", "\\immediate\\write18{ls}",
                 "\\csname input\\endcsname{/etc/hosts}"):
        assert compile_tex(JAKE.replace("\\begin{document}", "\\begin{document}" + evil)) is None


def test_parse_tex_upload_returns_style_and_source(client):
    r = client.post("/api/resume/parse", files={"file": ("cv.tex", JAKE.encode(), "application/x-tex")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["file_type"] == "tex" and body["source_id"]
    assert body["style"]["font_family"] == "serif" and body["style"]["heading_case"] == "smallcaps"
    assert body["style"]["margins"]["left"] == 0.5 and body["style"]["page_size"] == "Letter"
    assert body["layout"] == []


def test_plain_text_is_not_a_resume_but_is_a_supporting_doc(client):
    r = client.post("/api/resume/parse", files={"file": ("notes.txt", b"Built things " * 40, "text/plain")})
    assert r.status_code == 415 and "supporting" in r.json()["error"]["message"]
    r = client.post("/api/documents/parse", files={"file": ("notes.md", b"# QuickLink\nDeployed on Kubernetes (k3s) with Helm.", "text/markdown")})
    assert r.status_code == 200 and "Kubernetes" in r.json()["text"]
    r = client.post("/api/documents/parse", files={"file": ("cv.pdf", make_pdf(), "application/pdf")})
    assert r.status_code == 200 and r.json()["word_count"] > 100
    assert client.post("/api/documents/parse", files={"file": ("x.txt", b"hi", "text/plain")}).status_code == 422


def test_pdf_and_docx_style_detected(client):
    pdf = client.post("/api/resume/parse", files={"file": ("cv.pdf", make_pdf(), "application/pdf")}).json()
    assert pdf["style"]["source"] == "pdf" and pdf["style"]["font_family"] == "sans" and 0.25 <= pdf["style"]["margins"]["left"] <= 1.6
    dx = client.post("/api/resume/parse", files={"file": ("cv.docx", make_docx(), "application/octet-stream")}).json()
    assert dx["style"]["source"] == "docx" and dx["style"]["margins"]["left"] == 1.25


# --------------------------------------------------------------------------- #
def _tailor(client, text, **extra):
    r = client.post("/api/tailor", json={**_payload(text), **extra})
    assert r.status_code == 200, r.text
    return r.json()["resume"]


def test_original_docx_is_edited_in_place(client):
    import docx
    up = client.post("/api/resume/parse", files={"file": ("cv.docx", make_docx(), "application/octet-stream")}).json()
    resume = _tailor(client, up["text"])
    r = client.post("/api/export", json={"format": "docx", "resume": resume, "source_id": up["source_id"],
                                         "layout": {"template": "original"}})
    assert r.status_code == 200 and r.headers["x-render-mode"] == "original-file"
    d = docx.Document(io.BytesIO(r.content))
    assert "List Bullet" in {p.style.name for p in d.paragraphs}  # original list styling kept
    assert "AWS Cloud Practitioner" in [c.text for t in d.tables for row in t.rows for c in row.cells]  # table kept
    # A PDF of a DOCX original is a style match, and says so.
    r = client.post("/api/export", json={"format": "pdf", "resume": resume, "source_id": up["source_id"], "layout": {"template": "original"}})
    assert r.headers["x-render-mode"] == "style-match" and "Word download" in unquote(r.headers["x-render-note"])


def test_original_tex_export_and_inline_preview(client):
    up = client.post("/api/resume/parse", files={"file": ("cv.tex", JAKE.encode(), "application/x-tex")}).json()
    resume = _tailor(client, up["text"])
    r = client.post("/api/export", json={"format": "tex", "resume": resume, "source_id": up["source_id"], "layout": {"template": "original"}})
    assert r.headers["x-render-mode"] == "original-file" and r.text.startswith("%-------------------------")
    assert "\\section{Summary}" not in r.text  # summary is opt-in
    r = client.post("/api/export", json={"format": "pdf", "resume": resume, "source_id": up["source_id"],
                                         "layout": {"template": "original"}, "inline": True})
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and r.headers["content-disposition"].startswith("inline")
    assert r.headers["x-render-mode"] in ("original-file", "style-match")


def test_expired_source_falls_back_with_a_note(client):
    resume = _tailor(client, tex_to_text(JAKE)[0])
    r = client.post("/api/export", json={"format": "docx", "resume": resume, "source_id": "deadbeef", "layout": {"template": "original"}})
    assert r.status_code == 200 and r.headers["x-render-mode"] == "style-match" and "expired" in unquote(r.headers["x-render-note"])


def test_page_setup_changes_the_pdf(client):
    import pdfplumber
    resume = _tailor(client, tex_to_text(JAKE)[0])

    def first_x(layout):
        pdf = client.post("/api/export", json={"format": "pdf", "resume": resume, "layout": layout}).content
        with pdfplumber.open(io.BytesIO(pdf)) as p:
            page = p.pages[0]
            return min(c["x0"] for c in page.chars) / 72, page.width
    left_half, w = first_x({"margins": {"left": 0.5, "right": 0.5, "top": 0.5, "bottom": 0.5}, "page_size": "Letter"})
    left_one, _ = first_x({"margins": {"left": 1.0, "right": 1.0, "top": 1.0, "bottom": 1.0}})
    assert abs(left_half - 0.5) < 0.08 and abs(left_one - 1.0) < 0.08 and w == 612
    bad = client.post("/api/export", json={"format": "pdf", "resume": resume, "layout": {"margins": {"left": 5}}})
    assert bad.status_code == 422


def test_summary_toggle_in_exports(client):
    resume = _tailor(client, tex_to_text(JAKE)[0], sections=["summary", "skills", "experience", "projects"])
    assert resume["summary"]["tailored"]
    off = client.post("/api/export", json={"format": "txt", "resume": resume}).text
    on = client.post("/api/export", json={"format": "txt", "resume": resume, "layout": {"show_summary": True}}).text
    assert "SUMMARY" not in off and "SUMMARY" in on


def test_layout_settings_defaults():
    lay = LayoutSettings()
    assert lay.show_summary is False and lay.template == "classic" and lay.margins.left == 0.65
