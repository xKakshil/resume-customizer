"""Export a tailored resume / cover letter as an ATS-friendly PDF or DOCX.

Both formats use a single-column layout, real text (no images), standard
section headings and a common font — the things ATS parsers handle best.
"""

from __future__ import annotations

import io
import os
import re
from xml.sax.saxutils import escape

from app.models.schemas import TailoredResume
from app.services.render import _join

ACCENT = (0.13, 0.20, 0.35)
_FONT, _FONT_BOLD = "Helvetica", "Helvetica-Bold"
_FONT_PATHS = [("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")]


def _register_fonts() -> tuple[str, str]:
    """Use DejaVu (wide Unicode coverage) when present, falling back to Helvetica."""
    global _FONT, _FONT_BOLD
    if _FONT != "Helvetica":
        return _FONT, _FONT_BOLD
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_PATHS:
        if os.path.exists(regular) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont("Body", regular))
                pdfmetrics.registerFont(TTFont("Body-Bold", bold))
                from reportlab.lib.fonts import addMapping
                addMapping("Body", 0, 0, "Body")
                addMapping("Body", 1, 0, "Body-Bold")
                addMapping("Body", 0, 1, "Body")
                addMapping("Body", 1, 1, "Body-Bold")
                _FONT, _FONT_BOLD = "Body", "Body-Bold"
            except Exception:
                pass
            break
    return _FONT, _FONT_BOLD


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def resume_pdf(r: TailoredResume) -> bytes:
    """Render, and if the resume spills onto an extra page, retry with a tighter
    layout and keep it when that saves a page (one-page fit)."""
    data, pages = _resume_pdf(r, compact=False)
    if pages > 1:
        tight, tight_pages = _resume_pdf(r, compact=True)
        if tight_pages < pages:
            return tight
    return data


def pdf_page_count(data: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", data))


def _resume_pdf(r: TailoredResume, *, compact: bool) -> tuple[bytes, int]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    accent = colors.Color(*ACCENT)
    k = 0.93 if compact else 1.0  # type scale
    st = {
        "name": ParagraphStyle("name", fontName=bold, fontSize=19 * k, leading=23 * k, alignment=TA_CENTER, textColor=accent),
        "headline": ParagraphStyle("headline", fontName=font, fontSize=10.5 * k, leading=14 * k, alignment=TA_CENTER, textColor=colors.HexColor("#333333")),
        "contact": ParagraphStyle("contact", fontName=font, fontSize=8.8 * k, leading=12 * k, alignment=TA_CENTER, textColor=colors.HexColor("#444444")),
        "h": ParagraphStyle("h", fontName=bold, fontSize=10.5 * k, leading=13 * k, textColor=accent, spaceBefore=8 * k, spaceAfter=1),
        "body": ParagraphStyle("body", fontName=font, fontSize=9.6 * k, leading=12.6 * k),
        "role": ParagraphStyle("role", fontName=bold, fontSize=9.8 * k, leading=12.6 * k),
        "meta": ParagraphStyle("meta", fontName=font, fontSize=9 * k, leading=12.6 * k, alignment=2, textColor=colors.HexColor("#444444")),
        "sub": ParagraphStyle("sub", fontName=font, fontSize=9.2 * k, leading=12 * k, textColor=colors.HexColor("#444444")),
        "bullet": ParagraphStyle("bullet", fontName=font, fontSize=9.6 * k, leading=12.4 * k),
    }
    e = lambda s: escape(s or "")
    story: list = []
    width = A4[0] - 30 * mm

    if r.candidate.name:
        story.append(Paragraph(e(r.candidate.name), st["name"]))
    if r.candidate.headline:
        story.append(Paragraph(e(r.candidate.headline), st["headline"]))
    c = r.candidate.contact
    bits = [e(x) for x in (c.email, c.phone, c.location) if x]
    for l in c.links:
        if l.url:
            href = l.url if l.url.startswith(("http", "mailto")) else f"https://{l.url}"
            bits.append(f'<link href="{e(href)}" color="#1f3a68">{e(l.url.replace("https://", "").replace("http://", "").rstrip("/"))}</link>')
    if bits:
        story.append(Paragraph("  |  ".join(bits), st["contact"]))

    def heading(title: str):
        story.append(Paragraph(e(title.upper()), st["h"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=accent, spaceBefore=1, spaceAfter=4))

    st["li"] = ParagraphStyle("li", parent=st["bullet"], leftIndent=12, bulletIndent=2, bulletFontName=font, bulletFontSize=9, spaceAfter=1)

    def bullets(items: list[str]):
        # Real "•" characters (not drawn glyphs), so ATS parsers extract them as bullets.
        items = [i for i in items if i and i.strip()]
        if not items:
            return None
        return [Paragraph(e(i), st["li"], bulletText="•") for i in items]

    def entry(left: str, right: str, sub: str = "", items: list[str] | None = None):
        row = Table([[Paragraph(left, st["role"]), Paragraph(e(right), st["meta"])]], colWidths=[width * 0.7, width * 0.3])
        row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                 ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        block = [row]
        if sub:
            block.append(Paragraph(e(sub), st["sub"]))
        if items and (b := bullets(items)):
            block.extend(b)
        block.append(Spacer(1, 4))
        story.append(KeepTogether(block[:3]) if len(block) > 3 else KeepTogether(block))
        if len(block) > 3:
            story.extend(block[3:])

    if s := (r.summary.tailored or r.summary.original):
        heading("Summary")
        story.append(Paragraph(e(s), st["body"]))
    if r.skills.groups:
        heading("Skills")
        for g in r.skills.groups:
            label = f"<b>{e(g.category)}:</b> " if g.category else ""
            story.append(Paragraph(label + e(", ".join(g.items)), st["body"]))
    if r.experience:
        heading("Experience")
        for x in r.experience:
            left = e(x.title) + (f' <font name="{font}">— {e(x.company)}</font>' if x.company else "")
            entry(left, x.dates, x.location, [b.tailored or b.original for b in x.bullets])
    if r.projects:
        heading("Projects")
        for p in r.projects:
            left = e(p.name)
            if p.link:
                href = p.link if p.link.startswith("http") else f"https://{p.link}"
                left += f' <font name="{font}" size="8.5" color="#1f3a68"><link href="{e(href)}">[link]</link></font>'
            entry(left, p.dates, ", ".join(p.tech), [b.tailored or b.original for b in p.bullets])
    if r.education:
        heading("Education")
        for ed in r.education:
            entry(e(ed.degree) + (f' <font name="{font}">— {e(ed.institution)}</font>' if ed.institution else ""),
                  ed.dates, ed.location, ed.details)
    if r.certifications:
        heading("Certifications")
        if b := bullets(r.certifications):
            story.extend(b)
    for sec in r.additional_sections:
        if sec.items:
            heading(sec.title or "Additional")
            if b := bullets(sec.items):
                story.extend(b)

    buf = io.BytesIO()
    margin = (12 if compact else 15) * mm
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=(10 if compact else 13) * mm,
                            bottomMargin=(10 if compact else 13) * mm,
                            title=f"{r.candidate.name} — Resume".strip(" —"), author=r.candidate.name or "", creator="Resume Customizer")
    doc.build(story)
    return buf.getvalue(), doc.page


def letter_pdf(text: str, name: str = "") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font, _ = _register_fonts()
    style = ParagraphStyle("p", fontName=font, fontSize=11, leading=16)
    story = []
    for para in [p for p in text.replace("\r", "").split("\n\n")]:
        if para.strip():
            story += [Paragraph(escape(para.strip()).replace("\n", "<br/>"), style), Spacer(1, 9)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=25 * mm, bottomMargin=25 * mm,
                      title=f"Cover letter {name}".strip()).build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _docx_base():
    import docx
    from docx.shared import Mm, Pt

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(16)
        s.top_margin = s.bottom_margin = Mm(14)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    return d


def resume_docx(r: TailoredResume) -> bytes:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = _docx_base()
    accent = RGBColor(int(ACCENT[0] * 255), int(ACCENT[1] * 255), int(ACCENT[2] * 255))
    page_width = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin

    def para(text="", *, size=None, bold=False, color=None, align=None, after=0, before=0):
        p = d.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        return p

    def heading(title: str):
        p = para(title.upper(), size=11, bold=True, color=accent, before=9, after=3)
        pPr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "21335A")):
            bottom.set(qn(k), v)
        border.append(bottom)
        pPr.append(border)

    def bullet(text: str):
        p = d.add_paragraph(style="List Bullet")
        p.add_run(text)
        p.paragraph_format.space_after = Pt(1)

    def entry(left: str, company: str, right: str, sub: str, items: list[str]):
        p = para(before=3)
        p.paragraph_format.tab_stops.add_tab_stop(page_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        r1 = p.add_run(left)
        r1.bold = True
        if company:
            p.add_run(f" — {company}")
        if right:
            r3 = p.add_run(f"\t{right}")
            r3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        if sub:
            s = para(sub, size=9.5, color=RGBColor(0x44, 0x44, 0x44))
            s.paragraph_format.keep_with_next = True
        for i in items:
            if i:
                bullet(i)

    if r.candidate.name:
        para(r.candidate.name, size=20, bold=True, color=accent, align=WD_ALIGN_PARAGRAPH.CENTER)
    if r.candidate.headline:
        para(r.candidate.headline, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    c = r.candidate.contact
    if line := _join(c.email, c.phone, c.location, *[l.url for l in c.links], sep="  |  "):
        para(line, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

    if s := (r.summary.tailored or r.summary.original):
        heading("Summary")
        para(s)
    if r.skills.groups:
        heading("Skills")
        for g in r.skills.groups:
            p = para()
            if g.category:
                p.add_run(f"{g.category}: ").bold = True
            p.add_run(", ".join(g.items))
    if r.experience:
        heading("Experience")
        for x in r.experience:
            entry(x.title, x.company, x.dates, x.location, [b.tailored or b.original for b in x.bullets])
    if r.projects:
        heading("Projects")
        for pr in r.projects:
            entry(pr.name, "", pr.dates, _join(", ".join(pr.tech), pr.link), [b.tailored or b.original for b in pr.bullets])
    if r.education:
        heading("Education")
        for ed in r.education:
            entry(ed.degree, ed.institution, ed.dates, ed.location, ed.details)
    if r.certifications:
        heading("Certifications")
        for cert in r.certifications:
            bullet(cert)
    for sec in r.additional_sections:
        if sec.items:
            heading(sec.title or "Additional")
            for i in sec.items:
                bullet(i)

    d.core_properties.title = f"{r.candidate.name} Resume".strip()
    d.core_properties.author = r.candidate.name or ""
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def letter_docx(text: str, name: str = "") -> bytes:
    from docx.shared import Mm, Pt

    d = _docx_base()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(25)
        s.top_margin = s.bottom_margin = Mm(25)
    d.styles["Normal"].font.size = Pt(11)
    for para in text.replace("\r", "").split("\n\n"):
        if para.strip():
            p = d.add_paragraph(para.strip())
            p.paragraph_format.space_after = Pt(10)
    d.core_properties.title = f"Cover letter {name}".strip()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
