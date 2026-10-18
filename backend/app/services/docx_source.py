"""Write a tailored resume back into the user's own .docx, keeping its design.

Only the text of matched bullet paragraphs (and optionally the summary and
skills lines) changes; paragraph styles, numbering, fonts, spacing, tables,
headers and everything else in the file are left untouched. When bullets are
reordered, each paragraph *slot* keeps its formatting and receives the run
content of the bullet that moves into it.
"""

from __future__ import annotations

import copy
import io
import re
from difflib import SequenceMatcher

from docx.oxml.ns import qn

from app.services.parsing import _HEADING_RE

_BULLET_PREFIX = re.compile(r"^\s*([•●▪■◦‣⁃∙➢➤►▸✓✔\-–—*·])\s*")


def _norm(t: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", t.lower())


def _all_paragraphs(doc):
    out = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
    seen, uniq = set(), []
    for p in out:
        if id(p._p) not in seen:
            seen.add(id(p._p))
            uniq.append(p)
    return uniq


def _is_heading(text: str, kind: str | None = None) -> bool:
    t = text.strip().strip(":")
    if not t or len(t) > 45:
        return False
    if kind:
        return bool(_HEADING_RE[kind].match(t))
    return any(rx.match(t) for rx in _HEADING_RE.values())


def _content(p):
    return [el for el in p._p if el.tag != qn("w:pPr")]


def _replace_content(p, elements):
    for el in _content(p):
        p._p.remove(el)
    for el in elements:
        p._p.append(copy.deepcopy(el))


def _set_text(p, text: str):
    """Put `text` into the first text run (keeping its formatting); drop the rest."""
    runs = [el for el in _content(p) if el.tag in (qn("w:r"), qn("w:hyperlink"))]
    first = next((r for r in runs if r.tag == qn("w:r") and r.findall(qn("w:t"))), None)
    prefix = ""
    m = _BULLET_PREFIX.match(p.text or "")
    if m:
        prefix = m.group(1) + " "
    for r in runs:
        if r is not first:
            p._p.remove(r)
    if first is None:
        p.add_run(prefix + text)
        return
    ts = first.findall(qn("w:t"))
    ts[0].text = prefix + text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for t in ts[1:]:
        first.remove(t)


def _match(original: str, cands: list[tuple[int, str]], used: set[int]) -> int | None:
    key = _norm(original)
    if not key:
        return None
    best, ratio = None, 0.0
    for i, norm in cands:
        if i in used:
            continue
        if norm == key:
            return i
        r = SequenceMatcher(None, key, norm).ratio()
        if r > ratio:
            best, ratio = i, r
    return best if ratio >= 0.85 else None


def apply_to_docx(data: bytes, resume, show_summary: bool) -> tuple[bytes, dict]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    paras = _all_paragraphs(doc)
    norms = [(i, _norm(_BULLET_PREFIX.sub("", p.text or ""))) for i, p in enumerate(paras)]
    snapshots = {i: [copy.deepcopy(el) for el in _content(p)] for i, p in enumerate(paras)}
    used: set[int] = set()
    replaced = unmatched = 0

    for item in [*resume.experience, *resume.projects]:
        slots, moves = [], []
        for b in item.bullets:
            idx = _match(b.original, norms, used)
            if idx is None:
                unmatched += bool(b.original.strip())
                continue
            used.add(idx)
            slots.append(idx)
            moves.append((idx, b.tailored or b.original, b.original))
        for slot, (src, text, original) in zip(sorted(slots), moves):
            changed = text.strip() != original.strip()
            if not changed and slot == src:
                continue
            _replace_content(paras[slot], snapshots[src])
            if changed:
                _set_text(paras[slot], text)
            replaced += 1

    summary_action = _apply_summary(doc, resume, show_summary)
    skills_changed = _apply_skills(doc, resume)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), {"bullets_replaced": replaced, "bullets_unmatched": unmatched,
                            "summary": summary_action, "skills_lines": skills_changed}


def _section_body(paras, start: int) -> list:
    body = []
    for p in paras[start + 1:]:
        if _is_heading(p.text):
            break
        body.append(p)
    return body


def _apply_summary(doc, resume, show_summary: bool) -> str:
    paras = list(doc.paragraphs)
    text = (resume.summary.tailored or resume.summary.original or "").strip()
    head = next((i for i, p in enumerate(paras) if _is_heading(p.text, "summary")), None)
    if head is not None:
        body = [p for p in _section_body(paras, head) if p.text.strip()]
        if not show_summary:
            for p in [paras[head], *_section_body(paras, head)]:
                p._p.getparent().remove(p._p)
            return "removed"
        if text and body:
            _set_text(body[0], text)
            for p in body[1:]:
                p._p.getparent().remove(p._p)
            return "replaced"
        return "kept"
    if show_summary and text:
        first_head = next((i for i, p in enumerate(paras) if _is_heading(p.text)), None)
        if first_head is None:
            return "none"
        h = paras[first_head]
        new_head = copy.deepcopy(h._p)
        h._p.addprevious(new_head)
        from docx.text.paragraph import Paragraph
        hp = Paragraph(new_head, h._parent)
        _set_text(hp, "SUMMARY" if h.text.strip().isupper() else "Summary")
        template = next((p for p in _section_body(paras, first_head) if p.text.strip()), None)
        if template is not None:
            new_body = copy.deepcopy(template._p)
            ppr = new_body.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:numPr")) is not None:
                ppr.remove(ppr.find(qn("w:numPr")))
            new_head.addnext(new_body)
            _set_text(Paragraph(new_body, h._parent), text)
        return "added"
    return "none"


def _apply_skills(doc, resume) -> int:
    paras = _all_paragraphs(doc)
    head = next((i for i, p in enumerate(paras) if _is_heading(p.text, "skills")), None)
    if head is None:
        return 0
    changed = 0
    for p in _section_body(paras, head):
        txt = p.text or ""
        if ":" not in txt:
            continue
        label = txt.split(":", 1)[0].strip().lower()
        group = next((g for g in resume.skills.groups if g.category.strip().lower() == label), None)
        if not group or not group.items:
            continue
        new_items = ", ".join(group.items)
        if _norm(txt.split(":", 1)[1]) == _norm(new_items):
            continue
        runs = [el for el in _content(p) if el.tag == qn("w:r")]
        acc, colon_run = "", None
        for r in runs:
            acc += "".join(t.text or "" for t in r.findall(qn("w:t")))
            if ":" in acc:
                colon_run = r
                break
        if colon_run is None:
            continue
        ts = colon_run.findall(qn("w:t"))
        joined = "".join(t.text or "" for t in ts)
        head_part = joined[: joined.index(":") + 1] if ":" in joined else joined
        after = runs[runs.index(colon_run) + 1:]
        if after:
            ts[0].text = head_part
            for t in ts[1:]:
                colon_run.remove(t)
            at = after[0].findall(qn("w:t"))
            if at:
                at[0].text = " " + new_items
                at[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                for t in at[1:]:
                    after[0].remove(t)
            for r in after[1:]:
                p._p.remove(r)
        else:
            ts[0].text = head_part + " " + new_items
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            for t in ts[1:]:
                colon_run.remove(t)
        changed += 1
    return changed
